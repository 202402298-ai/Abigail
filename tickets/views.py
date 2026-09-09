import datetime
from collections import Counter

from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.shortcuts import redirect, render
from django.utils import timezone

from .forms import XMLUploadForm
from .models import AreaFavorita, ImportBatch, TeamArea, Ticket
from .services import metrics
from .services.hesk_parser import HeskParseError, parse_hesk_xml
from .services.importer import importar_tickets


def _areas_gestionadas(user):
    """Áreas que el usuario supervisa como 'jefe de área'. Lista vacía si el
    usuario no tiene ninguna asignada (ve todo, como un analista/consulta)."""
    return list(
        TeamArea.objects.filter(jefes__user=user).order_by('display_name')
    )


_DIAS_ABREV_ES = ['Lun', 'Mar', 'Mié', 'Jue', 'Vie', 'Sáb', 'Dom']


def _etiqueta_dia(fecha):
    """'2026-07-27' (lunes) -> 'Lun 27/07'."""
    return f'{_DIAS_ABREV_ES[fecha.weekday()]} {fecha.strftime("%d/%m")}'


def _tickets_visibles(user):
    """Ticket queryset que le corresponde ver a este usuario: si es jefe de
    una o más áreas, solo las suyas; si no, todo el sistema."""
    areas = _areas_gestionadas(user)
    if areas:
        codigos = [a.team_code for a in areas]
        return Ticket.objects.filter(team_code__in=codigos), areas
    return Ticket.objects.all(), areas


@login_required
def dashboard(request):
    tickets_qs, areas_gestionadas = _tickets_visibles(request.user)

    favoritas = list(TeamArea.objects.filter(favoritos__user=request.user))
    area_filtro = request.GET.get('area') or ''
    if area_filtro:
        tickets_qs = tickets_qs.filter(team_code=area_filtro)

    # Área específica que se está viendo, ya sea porque se eligió con el
    # filtro rápido de "Mis áreas" o porque el usuario es jefe de una sola
    # área: en ese caso no tiene sentido mostrar los reportes que comparan
    # entre todas las áreas.
    area_activa = area_filtro or (
        areas_gestionadas[0].team_code if len(areas_gestionadas) == 1 else ''
    )

    if not tickets_qs.exists():
        return render(request, 'tickets/dashboard.html', {
            'sin_datos': True, 'areas_gestionadas': areas_gestionadas,
            'favoritas': favoritas, 'area_filtro': area_filtro, 'area_activa': area_activa,
        })

    kpis = metrics.resumen(tickets_qs)
    por_area_pais, por_area_funcional = metrics.activos_por_area(tickets_qs)
    por_estado = metrics.tickets_por_estado(tickets_qs)
    por_prioridad = metrics.distribucion_prioridad(tickets_qs)
    buckets_tiempo_abierto, detalle_tiempo_abierto = metrics.tiempo_abierto(tickets_qs)
    _, tipo_area = metrics.tickets_por_categoria(tickets_qs)
    tiempos_global, tiempos_por_prioridad, tiempos_por_area = metrics.tiempos_de_servicio(tickets_qs)
    carga_agentes = metrics.carga_por_agente(tickets_qs)

    ranking_resolutores = metrics.resueltos_por_tecnico(area_activa) if area_activa else []
    semanas_periodo = metrics.DIAS_RANKING_RESOLUTORES / 7
    for fila in ranking_resolutores:
        fila['promedio_semanal'] = round(fila['total'] / semanas_periodo, 1)

    # Para ITSM, "quien cierra el ticket es quien le da seguimiento": la
    # gráfica muestra el promedio semanal en vez del total del mes (para
    # las demás áreas se deja el total, tal como estaba).
    if area_activa == 'ITSM':
        ranking_resolutores_chart = [r['promedio_semanal'] for r in ranking_resolutores]
    else:
        ranking_resolutores_chart = [r['total'] for r in ranking_resolutores]

    fechas_actividad, actividad_diaria = (
        metrics.actividad_diaria_itsm() if area_activa == 'ITSM' else ([], [])
    )

    # ITSM tiene su propio set de reportes (ver más abajo); "Todas" y el
    # resto de áreas comparten estos.
    mostrar_reportes_area = area_activa != 'ITSM'

    nuevos_por_dia = metrics.tickets_nuevos_por_dia(tickets_qs) if mostrar_reportes_area else []
    nuevos_por_dia_labels = [_etiqueta_dia(fila['fecha']) for fila in nuevos_por_dia]

    # "Tickets resueltos por día" solo tiene sentido para un área específica
    # (no para "Todas", que mezclaría resolutores de áreas distintas en un
    # mismo total sin desglose).
    resueltos_por_dia = (
        metrics.tickets_resueltos_por_dia(area_activa)
        if area_activa and area_activa != 'ITSM' else []
    )
    resueltos_por_dia_labels = [_etiqueta_dia(fila['fecha']) for fila in resueltos_por_dia]

    chart_data = {
        'por_area_pais': {
            'labels': [d['display_name'] for d in por_area_pais],
            'data': [d['total'] for d in por_area_pais],
        },
        'por_area_funcional': {
            'labels': [d['display_name'] for d in por_area_funcional],
            'data': [d['total'] for d in por_area_funcional],
        },
        'por_estado': {
            'labels': [d['status'] or '(sin estado)' for d in por_estado],
            'data': [d['total'] for d in por_estado],
        },
        'por_prioridad': {
            'labels': [d['priority'] or '(sin prioridad)' for d in por_prioridad],
            'data': [d['total'] for d in por_prioridad],
        },
        'tiempo_abierto': {
            'labels': list(buckets_tiempo_abierto.keys()),
            'data': list(buckets_tiempo_abierto.values()),
        },
        'ranking_resolutores': {
            'labels': [r['nombre'] for r in ranking_resolutores],
            'data': ranking_resolutores_chart,
        },
        'nuevos_por_dia': {
            'labels': nuevos_por_dia_labels,
            'data': [fila['total'] for fila in nuevos_por_dia],
        },
        'resueltos_por_dia': {
            'labels': resueltos_por_dia_labels,
            'data': [fila['total'] for fila in resueltos_por_dia],
        },
    }

    context = {
        'sin_datos': False,
        'kpis': kpis,
        'por_area_pais': por_area_pais,
        'por_area_funcional': por_area_funcional,
        'por_estado': por_estado,
        'por_prioridad': por_prioridad,
        'buckets_tiempo_abierto': buckets_tiempo_abierto,
        'detalle_tiempo_abierto': detalle_tiempo_abierto[:15],
        'tipo_area': tipo_area,
        'tiempos_global': tiempos_global,
        'tiempos_por_prioridad': tiempos_por_prioridad,
        'tiempos_por_area': tiempos_por_area,
        'carga_agentes': carga_agentes[:15],
        'mostrar_reportes_area': mostrar_reportes_area,
        'ranking_resolutores': ranking_resolutores,
        'fechas_actividad': fechas_actividad,
        'actividad_diaria': actividad_diaria,
        'chart_data': chart_data,
        'area_names': {ta.team_code: ta.display_name for ta in TeamArea.objects.all()},
        'ultima_importacion': ImportBatch.objects.first(),
        'areas_gestionadas': areas_gestionadas,
        'favoritas': favoritas,
        'area_filtro': area_filtro,
        'area_activa': area_activa,
    }
    return render(request, 'tickets/dashboard.html', context)


@login_required
def ticket_list(request):
    tickets_qs, areas_gestionadas = _tickets_visibles(request.user)
    tickets_qs = tickets_qs.select_related('last_batch')

    codes_presentes = (
        tickets_qs.values_list('team_code', flat=True).distinct().order_by('team_code')
    )
    area_lookup = {ta.team_code: ta.display_name for ta in TeamArea.objects.all()}
    team_code_options = [
        {'code': code, 'label': f"{area_lookup.get(code, code)} ({code})"}
        for code in codes_presentes
    ]

    team_code = request.GET.get('team_code')
    status = request.GET.get('status')
    priority = request.GET.get('priority')
    stage = request.GET.get('stage')
    last_batch = request.GET.get('last_batch')
    q = request.GET.get('q')

    if team_code:
        tickets_qs = tickets_qs.filter(team_code=team_code)
    if status:
        tickets_qs = tickets_qs.filter(status=status)
    if priority:
        tickets_qs = tickets_qs.filter(priority=priority)
    if stage:
        tickets_qs = tickets_qs.filter(stage=stage)
    if last_batch:
        tickets_qs = tickets_qs.filter(last_batch_id=last_batch)
    if q:
        tickets_qs = tickets_qs.filter(subject__icontains=q) | tickets_qs.filter(
            tracking_id__icontains=q
        )

    filtros = {
        'team_code_options': team_code_options,
        'statuses': Ticket.STATUS_CHOICES,
        'priorities': Ticket.PRIORITY_CHOICES,
        'stages': Ticket.STAGE_CHOICES,
        'batches': ImportBatch.objects.all()[:20],
    }

    context = {
        'tickets': tickets_qs.order_by('-created_at')[:500],
        'total_filtrado': tickets_qs.count(),
        'filtros': filtros,
        'area_names': area_lookup,
        'areas_gestionadas': areas_gestionadas,
        'selected': {
            'team_code': request.GET.get('team_code', ''),
            'status': request.GET.get('status', ''),
            'priority': request.GET.get('priority', ''),
            'stage': request.GET.get('stage', ''),
            'last_batch': request.GET.get('last_batch', ''),
            'q': request.GET.get('q', ''),
        },
    }
    return render(request, 'tickets/ticket_list.html', context)


@login_required
def mis_areas(request):
    if request.method == 'POST':
        codigos = request.POST.getlist('areas')
        AreaFavorita.objects.filter(user=request.user).delete()
        AreaFavorita.objects.bulk_create([
            AreaFavorita(user=request.user, team_area_id=codigo) for codigo in codigos
        ])
        messages.success(request, 'Tus áreas de seguimiento se actualizaron.')
        return redirect('tickets:mis_areas')

    seleccionadas = set(
        AreaFavorita.objects.filter(user=request.user).values_list('team_area_id', flat=True)
    )
    context = {
        'paises': TeamArea.objects.filter(area_type=TeamArea.TYPE_PAIS),
        'funcionales': TeamArea.objects.filter(area_type=TeamArea.TYPE_FUNCIONAL),
        'seleccionadas': seleccionadas,
    }
    return render(request, 'tickets/mis_areas.html', context)


DIAS_MINIMOS_SEGUIMIENTO = 3


@login_required
def seguimiento_detallado(request):
    if request.method == 'POST':
        actualizados = 0
        for clave, nota in request.POST.items():
            if not clave.startswith('nota_'):
                continue
            ticket_id = clave[len('nota_'):]
            actualizados += Ticket.objects.filter(pk=ticket_id).update(seguimiento_notas=nota)
        messages.success(request, 'Seguimiento guardado.')
        return redirect('tickets:seguimiento_detallado')

    tickets_qs, areas_gestionadas = _tickets_visibles(request.user)
    favoritas = list(TeamArea.objects.filter(favoritos__user=request.user))

    # Si el usuario no es jefe de área (ve todo) pero sí tiene áreas de
    # seguimiento configuradas, este módulo se enfoca solo en esas.
    if not areas_gestionadas and favoritas:
        tickets_qs = tickets_qs.filter(team_code__in=[a.team_code for a in favoritas])

    now = timezone.now()
    activos = (
        tickets_qs.exclude(status=Ticket.STATUS_RESUELTO)
        .exclude(created_at__isnull=True)
        .select_related('last_batch')
    )

    area_lookup = {ta.team_code: ta.display_name for ta in TeamArea.objects.all()}
    por_area = {}
    for t in activos:
        dias = (now - t.created_at).total_seconds() / 86400
        if dias <= DIAS_MINIMOS_SEGUIMIENTO:
            continue
        t.dias_abierto = int(dias)
        area_key = t.team_code or '(sin asignar)'
        tecnico_key = t.owner_name or '(sin asignar)'
        por_area.setdefault(area_key, {}).setdefault(tecnico_key, []).append(t)

    grupos = []
    for area_key in sorted(por_area, key=lambda c: area_lookup.get(c, c)):
        tecnicos = []
        total_area = 0
        for tecnico in sorted(por_area[area_key]):
            tickets_tecnico = sorted(por_area[area_key][tecnico], key=lambda t: t.dias_abierto, reverse=True)
            tecnicos.append({'nombre': tecnico, 'tickets': tickets_tecnico})
            total_area += len(tickets_tecnico)
        grupos.append({
            'area_code': area_key,
            'display_name': area_lookup.get(area_key, area_key),
            'tecnicos': tecnicos,
            'total': total_area,
        })

    context = {
        'grupos': grupos,
        'total_tickets': sum(g['total'] for g in grupos),
        'areas_gestionadas': areas_gestionadas,
        'favoritas': favoritas,
        'dias_minimos': DIAS_MINIMOS_SEGUIMIENTO,
    }
    return render(request, 'tickets/seguimiento_detallado.html', context)


def _rango_mes(request):
    """Lee ?mes=YYYY-MM de la query string (o el mes actual si no viene) y
    devuelve (inicio, fin, anio, mes) — fin es excluyente."""
    hoy = timezone.localdate()
    mes_param = request.GET.get('mes')
    try:
        anio, mes = (int(p) for p in mes_param.split('-')) if mes_param else (hoy.year, hoy.month)
    except ValueError:
        anio, mes = hoy.year, hoy.month

    inicio = timezone.make_aware(datetime.datetime(anio, mes, 1))
    fin = timezone.make_aware(
        datetime.datetime(anio + 1, 1, 1) if mes == 12 else datetime.datetime(anio, mes + 1, 1)
    )
    return inicio, fin, anio, mes


def _requiere_acceso_inf(user):
    areas_gestionadas = _areas_gestionadas(user)
    if areas_gestionadas and not any(a.team_code == 'INF' for a in areas_gestionadas):
        raise PermissionDenied


@login_required
def reporte_infraestructura(request):
    """Reporte que pidió el equipo de Infraestructura: tickets resueltos en
    un mes (con técnico responsable y colaboradores) + tickets pendientes
    con su estado actual."""
    _requiere_acceso_inf(request.user)

    hoy = timezone.localdate()
    inicio, fin, anio, mes = _rango_mes(request)

    resueltos = metrics.infraestructura_resueltos(inicio, fin)
    pendientes = metrics.infraestructura_pendientes()

    opciones_mes = []
    cursor = hoy.replace(day=1)
    for _ in range(12):
        opciones_mes.append(cursor)
        cursor = (cursor - datetime.timedelta(days=1)).replace(day=1)

    por_tecnico = Counter(f['ticket'].resuelto_por_nombre for f in resueltos)
    por_prioridad_resueltos = Counter(f['ticket'].priority or '(sin prioridad)' for f in resueltos)
    por_estado_pendientes = Counter(f['ticket'].status or '(sin estado)' for f in pendientes)

    por_dia = Counter()
    for f in resueltos:
        if f['ticket'].resuelto_por_fecha:
            por_dia[f['ticket'].resuelto_por_fecha.date()] += 1
    ultimo_dia = min(fin.date(), hoy + datetime.timedelta(days=1))
    dias_del_mes = []
    cursor_dia = inicio.date()
    while cursor_dia < ultimo_dia:
        dias_del_mes.append(cursor_dia)
        cursor_dia += datetime.timedelta(days=1)

    tecnicos_ordenados = por_tecnico.most_common(10)

    chart_data = {
        'por_tecnico': {
            'labels': [nombre for nombre, _ in tecnicos_ordenados],
            'data': [total for _, total in tecnicos_ordenados],
        },
        'por_prioridad_resueltos': {
            'labels': list(por_prioridad_resueltos.keys()),
            'data': list(por_prioridad_resueltos.values()),
        },
        'por_estado_pendientes': {
            'labels': list(por_estado_pendientes.keys()),
            'data': list(por_estado_pendientes.values()),
        },
        'por_dia': {
            'labels': [d.strftime('%d/%m') for d in dias_del_mes],
            'data': [por_dia.get(d, 0) for d in dias_del_mes],
        },
    }

    context = {
        'resueltos': resueltos,
        'pendientes': pendientes,
        'mes_actual': f'{anio:04d}-{mes:02d}',
        'opciones_mes': opciones_mes,
        'chart_data': chart_data,
    }
    return render(request, 'tickets/reporte_infraestructura.html', context)


def _hoja_excel(wb, nombre, encabezados, filas):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(nombre)
    ws.append(encabezados)
    for col, _ in enumerate(encabezados, start=1):
        celda = ws.cell(row=1, column=col)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='4C6EF5')
        celda.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'

    for fila in filas:
        ws.append(fila)

    anchos = [max(len(str(encabezados[i])), *(len(str(f[i])) if f[i] is not None else 0 for f in filas)) if filas else len(str(encabezados[i])) for i in range(len(encabezados))]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 10), 60)
    return ws


@login_required
def reporte_infraestructura_excel(request):
    """Descarga en Excel (.xlsx) las mismas dos tablas del reporte de
    Infraestructura, en hojas separadas, listas para filtrar/ordenar."""
    import openpyxl
    from django.http import HttpResponse

    _requiere_acceso_inf(request.user)

    inicio, fin, anio, mes = _rango_mes(request)
    resueltos = metrics.infraestructura_resueltos(inicio, fin)
    pendientes = metrics.infraestructura_pendientes()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _hoja_excel(
        wb, 'Resueltos',
        ['ID de ticket', 'ID de seguimiento', 'Fecha de creación', 'Fecha de resolución',
         'Técnico responsable', 'Técnico(s) colaborador(es)', 'Prioridad', 'Asunto', 'Descripción'],
        [
            [
                f['ticket'].hesk_row_id,
                f['ticket'].tracking_id,
                timezone.localtime(f['ticket'].created_at).strftime('%Y-%m-%d %H:%M') if f['ticket'].created_at else '',
                timezone.localtime(f['ticket'].resuelto_por_fecha).strftime('%Y-%m-%d %H:%M') if f['ticket'].resuelto_por_fecha else '',
                f['ticket'].resuelto_por_nombre,
                ', '.join(f['colaboradores']),
                f['ticket'].priority,
                f['ticket'].subject,
                f['ticket'].message,
            ]
            for f in resueltos
        ],
    )

    _hoja_excel(
        wb, 'Pendientes',
        ['ID de ticket', 'ID de seguimiento', 'Fecha de creación', 'Estado',
         'Técnico asignado', 'Prioridad', 'Asunto', 'Días abierto'],
        [
            [
                f['ticket'].hesk_row_id,
                f['ticket'].tracking_id,
                timezone.localtime(f['ticket'].created_at).strftime('%Y-%m-%d %H:%M') if f['ticket'].created_at else '',
                f['ticket'].status,
                f['ticket'].owner_name,
                f['ticket'].priority,
                f['ticket'].subject,
                f['dias_abierto'],
            ]
            for f in pendientes
        ],
    )

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    nombre_archivo = f'infraestructura_{anio:04d}-{mes:02d}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    wb.save(response)
    return response


@login_required
@permission_required('tickets.add_importbatch', raise_exception=True)
def upload_xml(request):
    if request.method == 'POST':
        form = XMLUploadForm(request.POST, request.FILES)
        if form.is_valid():
            xml_file = form.cleaned_data['xml_file']
            try:
                filas = parse_hesk_xml(xml_file)
            except HeskParseError as exc:
                messages.error(request, f'No se pudo leer el archivo: {exc}')
            else:
                if not filas:
                    messages.warning(request, 'El archivo no contiene tickets para importar.')
                else:
                    with transaction.atomic():
                        resultado = importar_tickets(filas, xml_file.name, request.user)
                    messages.success(
                        request,
                        f'Importación completa de "{xml_file.name}": '
                        f'{resultado.nuevos} tickets nuevos, {resultado.actualizados} actualizados '
                        f'y {resultado.sin_cambios} sin cambios (de {len(filas)} en el archivo).'
                    )
                    return redirect('tickets:dashboard')
    else:
        form = XMLUploadForm()

    historial = ImportBatch.objects.all()[:20]
    return render(request, 'tickets/upload.html', {'form': form, 'historial': historial})
